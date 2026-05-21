# -*- coding: utf-8 -*-
"""
insert_images.py
Substitui os placeholders [FIGURA X] no Relatorio_FitMap.docx pelas imagens reais.
Execute: python docs/insert_images.py
"""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image as PILImage
import io

HERE        = Path(__file__).parent
SCREENSHOTS = HERE / "screenshots"
DOCX_PATH   = HERE.parent / "Relatorio_FitMap.docx"

# Mapeamento: texto do placeholder -> (ficheiro, legenda)
FIGURA_MAP = {
    "FIGURA 1": (
        "01_landing.png",
        "Figura 1 - Landing page do FitMap: 3390 instalacoes, 9 categorias, 12+ cidades"
    ),
    "FIGURA 2": (
        "04_geonear_raio.png",
        "Figura 2 - Pesquisa por raio ($geoNear): instalacoes ordenadas por distancia, com auto-expansao 3->10 km"
    ),
    "FIGURA 3": (
        "03_lisboa_ginasios.png",
        "Figura 3 - Mapa de Lisboa filtrado por Ginasios (69 instalacoes visiveis na cidade)"
    ),
    "FIGURA 4": (
        "02_mapa_portugal.png",
        "Figura 4 - Vista nacional: mais de 3000 instalacoes georreferenciadas em todo o territorio"
    ),
    "FIGURA 5": (
        "06_detalhe_instalacao.png",
        "Figura 5 - Painel de detalhe: modalidades, acessibilidade e eventos proximos"
    ),
    "FIGURA 6": (
        "05_rota_eventos.png",
        "Figura 6 - Rota calculada via OSRM e painel de eventos desportivos proximos"
    ),
}


def make_image_paragraph(doc, img_path, width_cm=14):
    """Cria elemento XML de paragrafo com imagem centrada."""
    p_el = OxmlElement("w:p")

    # centrar
    pPr = OxmlElement("w:pPr")
    jc  = OxmlElement("w:jc")
    jc.set(qn("w:val"), "center")
    pPr.append(jc)
    p_el.append(pPr)

    # dimensoes
    EMU_PER_CM = 360000
    width_emu  = int(width_cm * EMU_PER_CM)
    with PILImage.open(str(img_path)) as im:
        w_px, h_px = im.size
    height_emu = int(width_emu * h_px / w_px)

    # adicionar imagem ao pacote docx e obter rId
    rId, _ = doc.part.get_or_add_image(str(img_path))

    # construir w:drawing
    drawing     = OxmlElement("w:drawing")
    inline      = OxmlElement("wp:inline")
    for attr in ("distT", "distB", "distL", "distR"):
        inline.set(attr, "0")

    extent = OxmlElement("wp:extent")
    extent.set("cx", str(width_emu))
    extent.set("cy", str(height_emu))
    inline.append(extent)

    docPr = OxmlElement("wp:docPr")
    docPr.set("id", str(abs(hash(str(img_path))) % 9000 + 1000))
    docPr.set("name", img_path.stem)
    inline.append(docPr)

    cNvGFP = OxmlElement("wp:cNvGraphicFramePr")
    gfl    = OxmlElement("a:graphicFrameLocks")
    gfl.set(qn("a:noChangeAspect"), "1")
    cNvGFP.append(gfl)
    inline.append(cNvGFP)

    graphic     = OxmlElement("a:graphic")
    graphicData = OxmlElement("a:graphicData")
    graphicData.set("uri", "http://schemas.openxmlformats.org/drawingml/2006/picture")

    pic     = OxmlElement("pic:pic")
    nvPicPr = OxmlElement("pic:nvPicPr")
    cNvPr2  = OxmlElement("pic:cNvPr")
    cNvPr2.set("id", "0")
    cNvPr2.set("name", img_path.name)
    nvPicPr.append(cNvPr2)
    nvPicPr.append(OxmlElement("pic:cNvPicPr"))
    pic.append(nvPicPr)

    blipFill = OxmlElement("pic:blipFill")
    blip     = OxmlElement("a:blip")
    blip.set(qn("r:embed"), rId)
    blipFill.append(blip)
    stretch  = OxmlElement("a:stretch")
    stretch.append(OxmlElement("a:fillRect"))
    blipFill.append(stretch)
    pic.append(blipFill)

    spPr  = OxmlElement("pic:spPr")
    xfrm  = OxmlElement("a:xfrm")
    off   = OxmlElement("a:off"); off.set("x","0"); off.set("y","0")
    ext2  = OxmlElement("a:ext"); ext2.set("cx", str(width_emu)); ext2.set("cy", str(height_emu))
    xfrm.append(off); xfrm.append(ext2)
    spPr.append(xfrm)
    prstGeom = OxmlElement("a:prstGeom"); prstGeom.set("prst","rect")
    prstGeom.append(OxmlElement("a:avLst"))
    spPr.append(prstGeom)
    pic.append(spPr)

    graphicData.append(pic)
    graphic.append(graphicData)
    inline.append(graphic)
    drawing.append(inline)

    r_el = OxmlElement("w:r")
    r_el.append(drawing)
    p_el.append(r_el)
    return p_el


def make_caption_paragraph(caption_text):
    """Cria paragrafo de legenda centrado e em italico."""
    p_el = OxmlElement("w:p")
    pPr  = OxmlElement("w:pPr")
    jc   = OxmlElement("w:jc"); jc.set(qn("w:val"), "center")
    spAfter = OxmlElement("w:spacing"); spAfter.set(qn("w:after"), "160")
    pPr.append(jc); pPr.append(spAfter)
    p_el.append(pPr)

    r_el = OxmlElement("w:r")
    rPr  = OxmlElement("w:rPr")
    rPr.append(OxmlElement("w:i"))
    sz   = OxmlElement("w:sz"); sz.set(qn("w:val"), "18")
    rPr.append(sz)
    r_el.append(rPr)
    t    = OxmlElement("w:t"); t.text = caption_text
    r_el.append(t)
    p_el.append(r_el)
    return p_el


def main():
    doc = Document(str(DOCX_PATH))
    body = doc.element.body

    # Encontrar paragrafos com [FIGURA X]
    replaced = 0
    paras = list(doc.paragraphs)
    for para in paras:
        text = para.text
        key  = None
        for k in FIGURA_MAP:
            if k in text:
                key = k
                break
        if key is None:
            continue

        img_file, caption = FIGURA_MAP[key]
        img_path = SCREENSHOTS / img_file

        if not img_path.exists():
            print(f"  [AVISO] Imagem nao encontrada: {img_file}")
            continue

        # Inserir imagem + legenda ANTES do paragrafo placeholder, depois apagar placeholder
        p_el   = para._element
        parent = p_el.getparent()
        idx    = list(parent).index(p_el)

        img_p = make_image_paragraph(doc, img_path)
        cap_p = make_caption_paragraph(caption)
        parent.insert(idx, cap_p)
        parent.insert(idx, img_p)
        parent.remove(p_el)   # remove o placeholder

        print(f"  OK  {key} -> {img_file}")
        replaced += 1

    doc.save(str(DOCX_PATH))
    print(f"\nConcluido: {replaced}/6 figuras inseridas -> {DOCX_PATH.name}")


if __name__ == "__main__":
    main()
